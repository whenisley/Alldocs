"""Generates the test documents in crates/alldocs/tests/fixtures.

Usage:
    pip install python-docx python-pptx openpyxl reportlab odfpy pillow
    python scripts/make_fixtures.py crates/alldocs/tests/fixtures
"""
import datetime
import io
import os
import sys
import tempfile
import zipfile

from PIL import Image, ImageDraw

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)


def png_bytes(color=(40, 120, 200), size=(120, 80), label="IMG"):
    img = Image.new("RGB", size, color)
    ImageDraw.Draw(img).text((10, 10), label, fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


PNG_PATH = os.path.join(tempfile.mkdtemp(), "_chart.png")
with open(PNG_PATH, "wb") as f:
    f.write(png_bytes())

# ---------------------------------------------------------------- DOCX
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


doc = Document()
doc.core_properties.title = "Quarterly Report"
doc.core_properties.author = "Ana Souza"
doc.add_heading("Quarterly Report", 0)
doc.add_heading("Summary", 1)
p = doc.add_paragraph("Revenue grew ")
p.add_run("12%").bold = True
p.add_run(" in ")
p.add_run("Q3").italic = True
p.add_run(", see the ")
add_hyperlink(p, "https://example.com/report", "full report")
p.add_run(" for details.")
p.add_run("[[FN]]")
doc.add_heading("Highlights", 2)
doc.add_paragraph("New markets opened", style="List Bullet")
doc.add_paragraph("Two regions", style="List Bullet 2")
doc.add_paragraph("Costs reduced", style="List Bullet")
doc.add_paragraph("Plan", style="List Number")
doc.add_paragraph("Execute", style="List Number")
doc.add_paragraph("Review", style="List Number")
doc.add_heading("Numbers", 2)
table = doc.add_table(rows=3, cols=3)
for r, row in enumerate([["Region", "Q2", "Q3"], ["North", "1,200", "1,350"], ["South | West", "900", "1,010"]]):
    for c, val in enumerate(row):
        table.cell(r, c).text = val
doc.add_heading("Code", 2)
for line in ["def total(xs):", "    return sum(xs)"]:
    cp = doc.add_paragraph()
    run = cp.add_run(line)
    run.font.name = "Courier New"
doc.add_paragraph("A wise quote about *growth* and_snake_case.", style="Quote")
doc.add_picture(PNG_PATH, width=Inches(1.5))
doc.add_paragraph("Strike ").add_run("this").font.strike = True
docx_path = os.path.join(OUT, "report.docx")
doc.save(docx_path)

# Inject a real footnote (python-docx has no footnote API).
FOOTNOTES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Unaudited figures.</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""
with zipfile.ZipFile(docx_path) as zin:
    entries = {n: zin.read(n) for n in zin.namelist()}
entries["word/footnotes.xml"] = FOOTNOTES.encode()
entries["word/document.xml"] = entries["word/document.xml"].replace(
    b"<w:t>[[FN]]</w:t>", b'<w:footnoteReference w:id="1"/>')
rels = entries["word/_rels/document.xml.rels"].decode()
rels = rels.replace("</Relationships>",
    '<Relationship Id="rIdFn1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
entries["word/_rels/document.xml.rels"] = rels.encode()
ct = entries["[Content_Types].xml"].decode().replace("</Types>",
    '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
entries["[Content_Types].xml"] = ct.encode()
with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
    for n, d in entries.items():
        zout.writestr(n, d)

# ---------------------------------------------------------------- PPTX
from pptx import Presentation
from pptx.util import Inches as PInches

prs = Presentation()
prs.core_properties.title = "Roadmap"
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Product Roadmap"
s.placeholders[1].text = "Planning for 2027"
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Goals"
tf = s.placeholders[1].text_frame
tf.text = "Grow revenue"
para = tf.add_paragraph(); para.text = "Enterprise segment"; para.level = 1
para = tf.add_paragraph(); para.text = "Improve quality"
run = tf.add_paragraph().add_run(); run.text = "Docs site"; run.hyperlink.address = "https://docs.example.com"
s.notes_slide.notes_text_frame.text = "Mention the Q3 numbers."
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Budget"
rows, cols = 3, 2
tbl = s.shapes.add_table(rows, cols, PInches(1), PInches(2), PInches(6), PInches(1.5)).table
for r, row in enumerate([["Item", "Cost"], ["Cloud", "$10k"], ["People", "$90k"]]):
    for c, v in enumerate(row):
        tbl.cell(r, c).text = v
s.shapes.add_picture(PNG_PATH, PInches(1), PInches(4))
prs.save(os.path.join(OUT, "roadmap.pptx"))

# ---------------------------------------------------------------- XLSX
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

wb = Workbook()
ws = wb.active
ws.title = "Sales"
ws.append(["Product", "Units", "Price", "Date", "Active"])
ws.append(["Widget", 10, 2.5, datetime.date(2026, 3, 1), True])
ws.append(["Gadget | Pro", 3, 0.1 + 0.2, datetime.datetime(2026, 3, 2, 14, 30), False])
ws.append([])
ws.append(["Totals", 13, None, None, None])
ws2 = wb.create_sheet("Notes")
ws2["B2"] = "Multi\nline"
ws2["C2"] = "x"
hidden = wb.create_sheet("Secret")
hidden["A1"] = "hidden data"
hidden.sheet_state = "hidden"
ws.add_image(XLImage(PNG_PATH), "G2")
wb.save(os.path.join(OUT, "sales.xlsx"))

# ---------------------------------------------------------------- ODT / ODS / ODP
from odf.opendocument import OpenDocumentText, OpenDocumentSpreadsheet, OpenDocumentPresentation
from odf.style import Style, TextProperties, ParagraphProperties, ListLevelProperties, PageLayout, MasterPage, PageLayoutProperties, GraphicProperties
from odf.text import H, P, Span, A, List, ListItem, Note, NoteCitation, NoteBody, ListStyle, ListLevelStyleNumber, ListLevelStyleBullet, LineBreak
from odf.table import Table, TableColumn, TableRow, TableCell
from odf.draw import Page, Frame, TextBox, Image as DrawImage
from odf import dc

odt = OpenDocumentText()
odt.meta.addElement(dc.Title(text="Field Notes"))
bold = Style(name="Bold", family="text"); bold.addElement(TextProperties(fontweight="bold"))
odt.automaticstyles.addElement(bold)
mono = Style(name="Mono", family="paragraph"); mono.addElement(TextProperties(fontfamily="Courier New", fontname="Courier New"))
odt.automaticstyles.addElement(mono)
numbered = ListStyle(name="Numbered")
numbered.addElement(ListLevelStyleNumber(level=1, numformat="1", numsuffix="."))
odt.automaticstyles.addElement(numbered)
bullets = ListStyle(name="Bullets")
bullets.addElement(ListLevelStyleBullet(level=1, bulletchar="•"))
bullets.addElement(ListLevelStyleBullet(level=2, bulletchar="◦"))
odt.automaticstyles.addElement(bullets)
odt.text.addElement(H(outlinelevel=1, text="Field Notes"))
p = P(text="Observed ")
p.addElement(Span(stylename=bold, text="three"))
p.addText(" species near the ")
p.addElement(A(href="https://example.org/river", text="river"))
p.addText(".")
note = Note(noteclass="footnote", id="ftn1")
note.addElement(NoteCitation(text="1"))
nb = NoteBody(); nb.addElement(P(text="Counted at dawn.")); note.addElement(nb)
p.addElement(note)
odt.text.addElement(p)
odt.text.addElement(H(outlinelevel=2, text="Steps"))
lst = List(stylename=numbered)
for t in ["Arrive", "Observe", "Record"]:
    li = ListItem(); li.addElement(P(text=t)); lst.addElement(li)
odt.text.addElement(lst)
bl = List(stylename=bullets)
li = ListItem(); li.addElement(P(text="Birds"))
sub = List(); sli = ListItem(); sli.addElement(P(text="Herons")); sub.addElement(sli); li.addElement(sub)
bl.addElement(li)
odt.text.addElement(bl)
t = Table(name="T1"); t.addElement(TableColumn(numbercolumnsrepeated=2))
for row in [["Species", "Count"], ["Heron", "4"]]:
    tr = TableRow()
    for v in row:
        tc = TableCell(); tc.addElement(P(text=v)); tr.addElement(tc)
    t.addElement(tr)
odt.text.addElement(t)
odt.text.addElement(P(stylename=mono, text="print('hello')"))
odt.save(os.path.join(OUT, "notes.odt"))

ods = OpenDocumentSpreadsheet()
t = Table(name="Inventory")
for row in [["Item", "Qty"], ["Apples", "5"], ["Pears", "7"]]:
    tr = TableRow()
    for v in row:
        tc = TableCell(valuetype="float", value=v) if v.isdigit() else TableCell(valuetype="string")
        tc.addElement(P(text=v)); tr.addElement(tc)
    t.addElement(tr)
ods.spreadsheet.addElement(t)
ods.save(os.path.join(OUT, "inventory.ods"))

odp = OpenDocumentPresentation()
pl = PageLayout(name="PL"); pl.addElement(PageLayoutProperties(pagewidth="28cm", pageheight="21cm"))
odp.automaticstyles.addElement(pl)
mp = MasterPage(name="Default", pagelayoutname=pl); odp.masterstyles.addElement(mp)
page = Page(name="page1", masterpagename=mp)
from odf.namespaces import PRESENTATIONNS
fr = Frame(width="20cm", height="3cm", x="2cm", y="1cm"); fr.setAttrNS(PRESENTATIONNS, "class", "title")
tb = TextBox(); tb.addElement(P(text="Kickoff")); fr.addElement(tb); page.addElement(fr)
fr = Frame(width="20cm", height="10cm", x="2cm", y="5cm"); fr.setAttrNS(PRESENTATIONNS, "class", "outline")
tb = TextBox(); ol = List()
for t_ in ["Scope", "Timeline"]:
    li = ListItem(); li.addElement(P(text=t_)); ol.addElement(li)
tb.addElement(ol); fr.addElement(tb); page.addElement(fr)
odp.presentation.addElement(page)
odp.save(os.path.join(OUT, "kickoff.odp"))

# ---------------------------------------------------------------- RTF
rtf = r"""{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\froman Times New Roman;}{\f1\fmodern\fprq1 Courier New;}}
{\stylesheet{\s0 Normal;}{\s1\outlinelevel0 heading 1;}{\s2\outlinelevel1 heading 2;}}
{\info{\title Meeting Minutes}{\author Rui}}
\pard\s1 Meeting Minutes\par
\pard Attendees discussed the \b budget\b0  and \i timeline\i0 . S\'e3o Paulo office UESC8212? confirmed.\par
\pard\s2 Actions\par
{\listtext 1.\tab}\pard\ls1\ilvl0 Draft plan\par
{\listtext 2.\tab}\pard\ls1\ilvl0 Send invites\par
\pard See {\field{\*\fldinst HYPERLINK "https://example.com/minutes"}{\fldrslt the archive}}.{\super\chftn}{\footnote\pard\plain {\super\chftn} Archive is internal.}\par
\trowd\cellx3000\cellx6000\pard\intbl Owner\cell Due\cell\row
\trowd\cellx3000\cellx6000\pard\intbl Rui\cell Friday\cell\row
\pard\f1 git push origin main\par
}"""
rtf = rtf.replace("UESC", chr(92) + "u")
with open(os.path.join(OUT, "minutes.rtf"), "w", encoding="ascii") as f:
    f.write(rtf)

# ---------------------------------------------------------------- EPUB
container = """<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>"""
opf = """<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>The Small Book</dc:title><dc:creator>J. Doe</dc:creator><dc:language>en</dc:language><dc:identifier id="id">x</dc:identifier></metadata>
<manifest>
<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>
<item id="c1" href="text/ch1.xhtml" media-type="application/xhtml+xml"/>
<item id="c2" href="text/ch2.xhtml" media-type="application/xhtml+xml"/>
<item id="img" href="images/fig1.png" media-type="image/png"/>
</manifest>
<spine><itemref idref="nav"/><itemref idref="c1"/><itemref idref="c2"/></spine></package>"""
nav = """<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><body><nav epub:type="toc"><ol><li><a href="text/ch1.xhtml">One</a></li></ol></nav></body></html>"""
ch1 = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>One</title></head><body>
<h1>Chapter One</h1>
<p>It was a <em>quiet</em> morning&#8212;the kind with &ldquo;fog&rdquo;.<a epub:type="noteref" href="#n1">1</a></p>
<p><img src="../images/fig1.png" alt="A figure"/></p>
<aside epub:type="footnote" id="n1"><p><a epub:type="backlink" href="#r1">1.</a> Fog is common here.</p></aside>
<pre><code class="language-python">print("hi")</code></pre>
</body></html>"""
ch2 = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Two</title></head><body>
<h1>Chapter Two</h1><p>See <a href="ch1.xhtml">chapter one</a> and <a href="https://example.com">the site</a>.</p>
<ol start="3"><li>three</li><li>four</li></ol>
<table><thead><tr><th>K</th><th>V</th></tr></thead><tbody><tr><td>a</td><td>1</td></tr></tbody></table>
</body></html>"""
with zipfile.ZipFile(os.path.join(OUT, "book.epub"), "w") as z:
    z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
    z.writestr("META-INF/container.xml", container)
    z.writestr("OEBPS/content.opf", opf)
    z.writestr("OEBPS/nav.xhtml", nav)
    z.writestr("OEBPS/text/ch1.xhtml", ch1)
    z.writestr("OEBPS/text/ch2.xhtml", ch2)
    z.writestr("OEBPS/images/fig1.png", png_bytes((200, 80, 40)))

# ---------------------------------------------------------------- CSV / TSV / HTML
with open(os.path.join(OUT, "people.csv"), "w", encoding="utf-8-sig", newline="") as f:
    f.write('Nome;Cidade;Nota\n"Silva, Ana";São Paulo;9,5\nJoão;"Rio\nde Janeiro";8\n')
with open(os.path.join(OUT, "data.tsv"), "w", encoding="utf-8", newline="") as f:
    f.write("a\tb\n1\t2\n")
with open(os.path.join(OUT, "page.html"), "w", encoding="utf-8") as f:
    f.write("<!DOCTYPE html><html lang=en><head><title>Page</title></head><body><h2>Hello</h2><p>Some <code>code</code> &amp; text</p></body></html>")

# ---------------------------------------------------------------- PDF
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

pdf_path = os.path.join(OUT, "paper.pdf")
c = canvas.Canvas(pdf_path, pagesize=A4)
c.setTitle("A Study of Things")
W, H = A4


def header_footer(n):
    c.setFont("Helvetica", 8)
    c.drawString(72, H - 40, "ACME Research - Internal")
    c.drawString(W / 2, 30, str(n))


def body_lines(y, lines, font="Helvetica", size=11, leading=14, x=72):
    c.setFont(font, size)
    for line in lines:
        c.drawString(x, y, line)
        y -= leading
    return y


header_footer(1)
c.setFont("Helvetica-Bold", 22)
c.drawString(72, H - 100, "A Study of Things")
c.setFont("Helvetica-Bold", 15)
c.drawString(72, H - 140, "1. Introduction")
y = body_lines(H - 165, [
    "This paper describes how things work in practice. We observed a large",
    "number of examples and found that most of them behave in a predict-",
    "able way under normal conditions.",
])
y -= 10
y = body_lines(y, ["A second paragraph starts here and continues for a little while longer",
                   "before it ends."])
y -= 10
y = body_lines(y, ["• First point", "• Second point", "• Third point"])
y -= 10
c.setFont("Helvetica-Bold", 15)
c.drawString(72, y - 10, "2. Method")
y = body_lines(y - 35, ["for x in things:", "    measure(x)"], font="Courier", size=10, leading=12)
c.drawImage(PNG_PATH, 72, y - 100, width=120, height=80)
c.showPage()
header_footer(2)
y = body_lines(H - 100, ["The method was applied to every item in the collection, and the results",
                         "were recorded carefully."])
c.showPage()
header_footer(3)
c.setFont("Helvetica-Bold", 15)
c.drawString(72, H - 100, "3. Results")
body_lines(H - 125, ["Results were good."])
c.showPage()
c.save()
print("fixtures written to", OUT)
